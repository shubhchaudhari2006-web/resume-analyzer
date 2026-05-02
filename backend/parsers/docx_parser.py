from docx import Document

class DOCXParser:
    """Parser for DOCX resume files"""
    
    @staticmethod
    def parse(file_path):
        """
        Extract text from DOCX file
        
        Args:
            file_path: Path to DOCX file
            
        Returns:
            Extracted text as string
        """
        try:
            doc = Document(file_path)
            text = ""
            
            # Extract paragraph text
            for para in doc.paragraphs:
                text += para.text + "\n"
            
            # Extract table text
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        text += cell.text + " "
                    text += "\n"
            
            return text.strip()
        except Exception as e:
            raise Exception(f"Error parsing DOCX: {str(e)}")
    
    @staticmethod
    def parse_from_bytes(file_bytes):
        """
        Extract text from DOCX bytes
        
        Args:
            file_bytes: DOCX file as bytes
            
        Returns:
            Extracted text as string
        """
        try:
            from io import BytesIO
            doc = Document(BytesIO(file_bytes))
            text = ""
            
            # Extract paragraph text
            for para in doc.paragraphs:
                text += para.text + "\n"
            
            # Extract table text
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        text += cell.text + " "
                    text += "\n"
            
            return text.strip()
        except Exception as e:
            raise Exception(f"Error parsing DOCX: {str(e)}")
